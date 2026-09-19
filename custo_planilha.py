"""Custo dos produtos vem de uma planilha de apoio mantida por fora do Tiny
(gerada por gerar_custos_por_codigo_pai.py e depois corrigida a mao pra
preencher os codigos sem custo encontrado), nao mais de consulta ao vivo na
API do Tiny. Motivo: o custo cadastrado no Tiny estava zerado/ausente numa
fatia grande dos produtos (ver README), e a planilha permite consolidar
custo de mais de uma fonte (Tiny + Mercado Livre) e corrigir manualmente o
que faltar sem depender de acesso de escrita ao Tiny.
"""

from __future__ import annotations

import logging
import os

import requests
from openpyxl import load_workbook

logger = logging.getLogger(__name__)

COL_CUSTO = 2
COL_SKUS_DO_GRUPO = 5


def carregar_custos(caminho: str) -> dict[str, float]:
    """Le custo_por_codigo_pai.xlsx e devolve {sku: custo}.

    A planilha e por codigo_pai (uma linha por grupo de variacoes), com a
    coluna skus_do_grupo listando todos os SKUs daquele grupo separados por
    virgula - aqui expandimos isso pra um dict direto por SKU, que e o que
    aparece em item["codigo"] no pedido do Tiny.
    """
    wb = load_workbook(caminho, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    custos: dict[str, float] = {}
    linhas = ws.iter_rows(values_only=True)
    next(linhas)  # cabecalho: codigo_pai, descricao, custo, fonte, conflito, skus_do_grupo

    for row in linhas:
        custo = row[COL_CUSTO]
        skus_do_grupo = row[COL_SKUS_DO_GRUPO]
        if custo in (None, "") or not skus_do_grupo:
            continue
        for sku in str(skus_do_grupo).split(","):
            sku = sku.strip()
            if sku:
                custos[sku] = float(custo)

    return custos


_PAGINA = 1000


def sobrepor_custos_do_dashboard(custos: dict[str, float]) -> dict[str, float]:
    """Sobrepoe (in-place, e retorna) `custos` com os valores que estao em
    margin_monitor_custos no Supabase - tanto os migrados da planilha
    original (gerar_custos_para_supabase.py) quanto os que o time preencheu
    direto no dashboard. Esses ganham prioridade sobre a planilha local,
    ja que sao a fonte mais recente.

    Passar `custos={}` carrega SOMENTE do Supabase - e o que o
    monitor_cloud.py faz (sem planilha local disponivel no runner).

    Pagina em blocos de 1000 (limite padrao do Supabase) - a tabela ja
    passou de ~2000 linhas depois da migracao da planilha original.

    Falha de rede/credencial aqui NUNCA deve impedir o monitor de rodar -
    so loga e segue com o que ja tinha em `custos`.
    """
    try:
        url = os.environ["SUPABASE_URL"].rstrip("/")
        key = os.environ["SUPABASE_SERVICE_ROLE_KEY"]
        headers = {"apikey": key, "Authorization": f"Bearer {key}"}

        inicio = 0
        while True:
            resp = requests.get(
                f"{url}/rest/v1/margin_monitor_custos",
                params={"select": "sku,custo"},
                headers={**headers, "Range": f"{inicio}-{inicio + _PAGINA - 1}"},
                timeout=15,
            )
            resp.raise_for_status()
            pagina = resp.json()
            for linha in pagina:
                custos[linha["sku"]] = float(linha["custo"])
            if len(pagina) < _PAGINA:
                break
            inicio += _PAGINA
    except (KeyError, requests.RequestException):
        logger.warning("Nao foi possivel buscar custos do Supabase - usando so o que ja tinha")

    return custos


def recalcular_custo_pendente() -> int:
    """Recalcula margem dos itens/pedidos ja gravados com custo_ausente=True
    cujo SKU ganhou custo depois (preenchido pelo time no dashboard, na aba
    "Custos pendentes" - o custo "zerado" foi corrigido).

    So pode rodar aqui (monitor.py/monitor_cloud.py), nunca a partir do
    dashboard: margin_monitor_itens/pedidos so aceitam escrita da service
    role key por RLS - o dashboard publico so tem a chave anon, que so pode
    escrever em margin_monitor_custos (tabela de baixo risco). E por isso
    que o preenchimento no dashboard so atualiza o cadastro na hora; o
    historico e corrigido aqui, no proximo ciclo do monitor (ate 15 min de
    atraso no GitHub Actions).

    custo_unitario era None nesses itens (cmv contribuido = 0), entao o novo
    cmv e integralmente descontado da margem ja salva - nao precisa
    reconstruir imposto/comissao/frete/ads de cada canal.

    Retorna quantos itens foram corrigidos. Falha de rede/credencial aqui
    NUNCA deve impedir o ciclo do monitor de continuar.
    """
    try:
        url = os.environ["SUPABASE_URL"].rstrip("/")
        key = os.environ["SUPABASE_SERVICE_ROLE_KEY"]
    except KeyError:
        return 0

    headers = {"apikey": key, "Authorization": f"Bearer {key}", "Content-Type": "application/json"}

    try:
        # O Supabase tem um teto de 1000 linhas por resposta (db-max-rows) -
        # ignora `limit` maior que isso. Pagina de verdade via Range (bug
        # real, achado em 2026-09-19 no carregar_tabela do dashboard.py -
        # mesma causa aqui: silenciosamente so pegava as 1000 primeiras
        # linhas de custo_ausente=true, deixando o resto pra tras).
        pendentes: list[dict] = []
        inicio = 0
        while True:
            resp = requests.get(
                f"{url}/rest/v1/margin_monitor_itens",
                params={"select": "conta_tiny,numero_pedido,sku", "custo_ausente": "eq.true"},
                headers={**headers, "Range": f"{inicio}-{inicio + _PAGINA - 1}"},
                timeout=20,
            )
            resp.raise_for_status()
            pagina = resp.json()
            pendentes.extend(pagina)
            if len(pagina) < _PAGINA:
                break
            inicio += _PAGINA
        if not pendentes:
            return 0

        skus_pendentes = sorted({p["sku"] for p in pendentes})
        resp = requests.get(
            f"{url}/rest/v1/margin_monitor_custos",
            params={"select": "sku,custo", "sku": f"in.({','.join(skus_pendentes)})"},
            headers=headers,
            timeout=20,
        )
        resp.raise_for_status()
        custo_por_sku = {linha["sku"]: float(linha["custo"]) for linha in resp.json() if linha["custo"]}
        if not custo_por_sku:
            return 0

        pedidos_afetados = sorted(
            {(p["conta_tiny"], p["numero_pedido"]) for p in pendentes if p["sku"] in custo_por_sku}
        )
        if not pedidos_afetados:
            return 0

        total_itens_corrigidos = 0
        for conta_tiny, numero_pedido in pedidos_afetados:
            resp = requests.get(
                f"{url}/rest/v1/margin_monitor_itens",
                params={
                    "select": "conta_tiny,numero_pedido,sku,canal,data_pedido,quantidade,"
                    "receita,cmv,margem_contribuicao,margem_pct,custo_ausente",
                    "conta_tiny": f"eq.{conta_tiny}",
                    "numero_pedido": f"eq.{numero_pedido}",
                },
                headers=headers,
                timeout=20,
            )
            resp.raise_for_status()
            itens_pedido = resp.json()

            itens_corrigidos = []
            for item in itens_pedido:
                if item["custo_ausente"] and item["sku"] in custo_por_sku:
                    cmv_novo = round(item["quantidade"] * custo_por_sku[item["sku"]], 2)
                    margem_novo = round(item["margem_contribuicao"] - cmv_novo, 2)
                    receita_item = item["receita"]
                    item = {
                        **item,
                        "cmv": cmv_novo,
                        "margem_contribuicao": margem_novo,
                        "margem_pct": round(margem_novo / receita_item * 100, 2) if receita_item else 0.0,
                        "custo_ausente": False,
                    }
                    total_itens_corrigidos += 1
                itens_corrigidos.append(item)

            resp = requests.post(
                f"{url}/rest/v1/margin_monitor_itens?on_conflict=conta_tiny,numero_pedido,sku",
                json=itens_corrigidos,
                headers={**headers, "Prefer": "resolution=merge-duplicates,return=minimal"},
                timeout=20,
            )
            resp.raise_for_status()

            margem_pedido = round(sum(it["margem_contribuicao"] for it in itens_corrigidos), 2)
            receita_pedido = round(sum(it["receita"] for it in itens_corrigidos), 2)
            custo_ausente_pedido = any(it["custo_ausente"] for it in itens_corrigidos)

            resp = requests.patch(
                f"{url}/rest/v1/margin_monitor_pedidos",
                params={"conta_tiny": f"eq.{conta_tiny}", "numero_pedido": f"eq.{numero_pedido}"},
                json={
                    "margem_contribuicao": margem_pedido,
                    "margem_pct": round(margem_pedido / receita_pedido * 100, 2) if receita_pedido else 0.0,
                    "custo_ausente": custo_ausente_pedido,
                },
                headers={**headers, "Prefer": "return=minimal"},
                timeout=20,
            )
            resp.raise_for_status()

        return total_itens_corrigidos
    except requests.RequestException:
        logger.exception("Falha ao recalcular custo pendente - tentando de novo no proximo ciclo")
        return 0
