"""Gera uma planilha de apoio (custo por codigo pai) cruzando 2 exports:

1. Export do cadastro do Tiny (aba "BaseDados") - traz Codigo (SKU),
   Preco de custo e Codigo do pai.
2. Export de publicacoes do Mercado Livre (aba "Publicacoes") - traz SKU e
   "Custo do produto carregado" (custo cadastrado direto no ML).

Agrupa por codigo pai (varias variacoes de tamanho/cor compartilham o mesmo
pai) e usa qualquer custo nao-zero encontrado em qualquer SKU do grupo, em
qualquer uma das 2 fontes. Sinaliza codigo pai sem custo em nenhuma fonte, e
codigo pai com custos diferentes conflitantes entre SKUs do mesmo grupo.

Uso:
    python gerar_custos_por_codigo_pai.py --base-dados "<export do Tiny>.xlsx" --custo-ml "<export do ML>.xlsx"
"""

from __future__ import annotations

import argparse
from collections import defaultdict

from openpyxl import Workbook, load_workbook

COL_BASE_SKU = 1
COL_BASE_PRECO_CUSTO = 11
COL_BASE_CODIGO_PAI = 37
COL_BASE_DESCRICAO = 2

COL_ML_SKU = 2
COL_ML_CUSTO_CARREGADO = 7


def _custo_valido(valor) -> float | None:
    try:
        valor = float(valor)
    except (TypeError, ValueError):
        return None
    return valor if valor > 0 else None


def carregar_base_dados(caminho: str) -> dict[str, dict]:
    """Retorna {sku: {"codigo_pai": ..., "custo_tiny": ..., "descricao": ...}}"""
    wb = load_workbook(caminho, read_only=True, data_only=True)
    ws = wb["BaseDados"]

    produtos: dict[str, dict] = {}
    linhas = ws.iter_rows(values_only=True)
    next(linhas)  # cabecalho

    for row in linhas:
        sku = row[COL_BASE_SKU]
        if not sku:
            continue
        sku = str(sku).strip()
        codigo_pai = row[COL_BASE_CODIGO_PAI]
        codigo_pai = str(codigo_pai).strip() if codigo_pai else sku  # sem pai cadastrado = pai de si mesmo
        produtos[sku] = {
            "codigo_pai": codigo_pai,
            "custo_tiny": _custo_valido(row[COL_BASE_PRECO_CUSTO]),
            "descricao": row[COL_BASE_DESCRICAO] or "",
        }
    return produtos


def carregar_custo_ml(caminho: str) -> dict[str, float]:
    """Retorna {sku: custo_carregado_no_ml}"""
    wb = load_workbook(caminho, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    custos: dict[str, float] = {}
    linhas = ws.iter_rows(values_only=True)
    next(linhas)  # cabecalho

    for row in linhas:
        sku = row[COL_ML_SKU]
        if not sku:
            continue
        sku = str(sku).strip()
        custo = _custo_valido(row[COL_ML_CUSTO_CARREGADO])
        if custo is not None:
            custos[sku] = custo
    return custos


def montar_custo_por_codigo_pai(produtos: dict[str, dict], custos_ml: dict[str, float]) -> dict[str, dict]:
    """Agrupa por codigo_pai e resolve um custo consolidado por grupo."""
    grupos: dict[str, list[dict]] = defaultdict(list)

    for sku, dados in produtos.items():
        custo_tiny = dados["custo_tiny"]
        custo_ml = custos_ml.get(sku)
        grupos[dados["codigo_pai"]].append(
            {
                "sku": sku,
                "descricao": dados["descricao"],
                "custo_tiny": custo_tiny,
                "custo_ml": custo_ml,
            }
        )

    resultado: dict[str, dict] = {}
    for codigo_pai, itens in grupos.items():
        custos_encontrados = set()
        fontes = set()
        for item in itens:
            if item["custo_tiny"] is not None:
                custos_encontrados.add(round(item["custo_tiny"], 2))
                fontes.add("Tiny")
            if item["custo_ml"] is not None:
                custos_encontrados.add(round(item["custo_ml"], 2))
                fontes.add("ML")

        if not custos_encontrados:
            resultado[codigo_pai] = {
                "custo": None,
                "fonte": "NENHUMA",
                "conflito": False,
                "skus": [i["sku"] for i in itens],
                "descricao": itens[0]["descricao"],
            }
        else:
            custo_escolhido = sorted(custos_encontrados)[0] if len(custos_encontrados) == 1 else max(custos_encontrados)
            resultado[codigo_pai] = {
                "custo": custo_escolhido,
                "fonte": "+".join(sorted(fontes)),
                "conflito": len(custos_encontrados) > 1,
                "skus": [i["sku"] for i in itens],
                "descricao": itens[0]["descricao"],
            }

    return resultado


def salvar_planilha(resultado: dict[str, dict], caminho_saida: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Custo por codigo pai"
    ws.append(["codigo_pai", "descricao", "custo", "fonte", "conflito", "skus_do_grupo"])

    for codigo_pai, dados in sorted(resultado.items()):
        ws.append(
            [
                codigo_pai,
                dados["descricao"],
                dados["custo"] if dados["custo"] is not None else "",
                dados["fonte"],
                "SIM" if dados["conflito"] else "",
                ", ".join(dados["skus"]),
            ]
        )

    wb.save(caminho_saida)


def _main():
    parser = argparse.ArgumentParser(description="Gera planilha de custo por codigo pai")
    parser.add_argument("--base-dados", required=True, help="Export do Tiny (aba BaseDados)")
    parser.add_argument("--custo-ml", required=True, help="Export de publicacoes do ML (custo carregado)")
    parser.add_argument("--saida", default="custo_por_codigo_pai.xlsx")
    args = parser.parse_args()

    produtos = carregar_base_dados(args.base_dados)
    custos_ml = carregar_custo_ml(args.custo_ml)
    resultado = montar_custo_por_codigo_pai(produtos, custos_ml)
    salvar_planilha(resultado, args.saida)

    sem_custo = sorted(cp for cp, d in resultado.items() if d["custo"] is None)
    conflitos = sorted(cp for cp, d in resultado.items() if d["conflito"])

    print(f"{len(resultado)} codigos pai no total")
    print(f"{len(sem_custo)} SEM custo em nenhuma fonte")
    print(f"{len(conflitos)} com custo CONFLITANTE entre SKUs do mesmo grupo")
    print(f"Planilha salva em {args.saida}")

    if sem_custo:
        print("\nCodigos pai sem custo (amostra ate 30):")
        for cp in sem_custo[:30]:
            print(" -", cp, "|", resultado[cp]["descricao"][:60])


if __name__ == "__main__":
    _main()
