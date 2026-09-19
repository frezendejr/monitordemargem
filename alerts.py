"""Envio de alertas (WhatsApp, e-mail) e atualizacao da planilha de historico."""

from __future__ import annotations

import logging
import os
import smtplib
from email.mime.text import MIMEText
from pathlib import Path

import requests
from openpyxl import Workbook, load_workbook

from margin_engine import ResultadoMargem

logger = logging.getLogger(__name__)

XLSX_PATH_PADRAO = "dashboard_margem.xlsx"
_CABECALHO = [
    "numero_pedido",
    "canal",
    "receita",
    "cmv",
    "imposto",
    "comissao",
    "frete",
    "ads",
    "margem_contribuicao",
    "margem_pct",
    "custo_ausente",
    "skus_sem_custo",
]


def montar_mensagem_alerta(resultado: ResultadoMargem) -> str:
    linhas = [
        "⚠️ Margem negativa detectada",
        f"Pedido: {resultado.numero_pedido}",
        f"Canal: {resultado.canal}",
        f"Receita: R$ {resultado.receita:.2f}",
        f"Margem de contribuicao: R$ {resultado.margem_contribuicao:.2f} ({resultado.margem_pct:.1f}%)",
    ]
    if resultado.custo_ausente:
        linhas.append(
            f"❗ Custo ausente no Tiny para: {', '.join(resultado.skus_sem_custo)} "
            "(margem calculada pode estar otimista - corrigir cadastro)"
        )
    return "\n".join(linhas)


# ----------------------------------------------------------------------
# WhatsApp
# ----------------------------------------------------------------------

def enviar_whatsapp(mensagem: str, config: dict) -> None:
    wa_config = config.get("alertas", {}).get("whatsapp", {})
    if not wa_config.get("ativo"):
        return

    provider = wa_config.get("provider", "zapi")
    destinatarios = wa_config.get("destinatarios", [])

    for numero in destinatarios:
        try:
            if provider == "zapi":
                _enviar_whatsapp_zapi(numero, mensagem)
            elif provider == "meta":
                _enviar_whatsapp_meta(numero, mensagem)
            elif provider == "callmebot":
                _enviar_whatsapp_callmebot(numero, mensagem)
            else:
                logger.error("Provider de WhatsApp desconhecido: %s", provider)
        except KeyError as e:
            logger.error(
                "WhatsApp ativo=true no config mas falta credencial %s no .env - alerta nao enviado", e
            )
        except requests.RequestException:
            logger.exception("Falha ao enviar WhatsApp para %s via %s", numero, provider)


def _enviar_whatsapp_zapi(numero: str, mensagem: str) -> None:
    instance_id = os.environ["ZAPI_INSTANCE_ID"]
    token = os.environ["ZAPI_TOKEN"]
    client_token = os.environ.get("ZAPI_CLIENT_TOKEN", "")

    url = f"https://api.z-api.io/instances/{instance_id}/token/{token}/send-text"
    headers = {"Client-Token": client_token} if client_token else {}
    resp = requests.post(
        url,
        json={"phone": numero, "message": mensagem},
        headers=headers,
        timeout=15,
    )
    resp.raise_for_status()


def _enviar_whatsapp_meta(numero: str, mensagem: str) -> None:
    token = os.environ["META_WHATSAPP_TOKEN"]
    phone_number_id = os.environ["META_WHATSAPP_PHONE_NUMBER_ID"]

    url = f"https://graph.facebook.com/v20.0/{phone_number_id}/messages"
    headers = {"Authorization": f"Bearer {token}"}
    payload = {
        "messaging_product": "whatsapp",
        "to": numero,
        "type": "text",
        "text": {"body": mensagem},
    }
    resp = requests.post(url, json=payload, headers=headers, timeout=15)
    resp.raise_for_status()


def _enviar_whatsapp_callmebot(numero: str, mensagem: str) -> None:
    """CallMeBot - gratuito, mas cada numero precisa ter ativado o bot
    individualmente (mandar "I allow callmebot to send me messages" pro
    contato +34 621 47 20 71) e tem sua PROPRIA apikey - por isso a chave no
    .env e por numero (CALLMEBOT_APIKEY_<numero>), nao uma unica global."""
    apikey = os.environ[f"CALLMEBOT_APIKEY_{numero}"]
    resp = requests.get(
        "https://api.callmebot.com/whatsapp.php",
        params={"phone": numero, "text": mensagem, "apikey": apikey},
        timeout=15,
    )
    resp.raise_for_status()


# ----------------------------------------------------------------------
# E-mail
# ----------------------------------------------------------------------

def enviar_email(assunto: str, corpo: str, config: dict) -> None:
    email_config = config.get("alertas", {}).get("email", {})
    if not email_config.get("ativo"):
        return

    destinatarios = email_config.get("destinatarios", [])
    if not destinatarios:
        return

    host = os.environ.get("SMTP_HOST", "")
    usuario = os.environ.get("SMTP_USER", "")
    senha = os.environ.get("SMTP_PASSWORD", "")
    if not host or not usuario or not senha:
        # Antes so pegava KeyError (variavel ausente) - mas SMTP_HOST/USER/
        # PASSWORD existem no .env como string vazia (SMTP ainda nao
        # configurado), entao passava direto pro smtplib.SMTP("", ...) e
        # falhava la dentro com um traceback confuso
        # (SMTPServerDisconnected: please run connect() first) a cada alerta.
        logger.error("E-mail ativo=true no config mas SMTP_HOST/USER/PASSWORD nao preenchidos no .env - alerta nao enviado")
        return

    port = int(os.environ.get("SMTP_PORT", 587))
    remetente = os.environ.get("SMTP_REMETENTE", usuario)

    msg = MIMEText(corpo, "plain", "utf-8")
    msg["Subject"] = assunto
    msg["From"] = remetente
    msg["To"] = ", ".join(destinatarios)

    try:
        with smtplib.SMTP(host, port, timeout=15) as smtp:
            smtp.starttls()
            smtp.login(usuario, senha)
            smtp.sendmail(remetente, destinatarios, msg.as_string())
    except smtplib.SMTPException:
        logger.exception("Falha ao enviar e-mail de alerta")


# ----------------------------------------------------------------------
# Planilha de historico
# ----------------------------------------------------------------------

def atualizar_planilha(resultado: ResultadoMargem, caminho: str = XLSX_PATH_PADRAO) -> None:
    path = Path(caminho)

    if path.exists():
        wb = load_workbook(path)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Historico"
        ws.append(_CABECALHO)

    ws.append(
        [
            resultado.numero_pedido,
            resultado.canal,
            resultado.receita,
            resultado.cmv,
            resultado.imposto,
            resultado.comissao,
            resultado.frete,
            resultado.ads,
            resultado.margem_contribuicao,
            resultado.margem_pct,
            "SIM" if resultado.custo_ausente else "",
            ", ".join(resultado.skus_sem_custo),
        ]
    )
    wb.save(path)
